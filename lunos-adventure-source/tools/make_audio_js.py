# audio.js（BGM・効果音の設定）を書き出す。引数に xlsx を渡すと、シートの内容で作る（なければ空の設定）
import sys, json, os
sys.path.insert(0, os.path.dirname(__file__)); from audio_def import BGM, SE
cfg = {'BGM': {k: {'file': '', 'vol': 60, 'loop': True, 'fade': 1.0} for k, _ in BGM}, 'SE': {k: {'file': '', 'vol': 80} for k, _, _ in SE}}
if len(sys.argv) > 2:
    from openpyxl import load_workbook
    wb = load_workbook(sys.argv[2], data_only=True)
    for r in wb['BGM'].iter_rows(min_row=2, values_only=True):
        if r[1] in cfg['BGM']:
            e = cfg['BGM'][r[1]]; e['file'] = (str(r[3]).strip() if r[3] else ''); e['vol'] = int(r[4]) if r[4] not in (None, '') else 60
            e['loop'] = (str(r[5]).strip() != 'しない'); e['fade'] = float(r[6]) if r[6] not in (None, '') else 1.0
    for r in wb['効果音'].iter_rows(min_row=2, values_only=True):
        if r[1] in cfg['SE']:
            e = cfg['SE'][r[1]]; e['file'] = (str(r[4]).strip() if r[4] else ''); e['vol'] = int(r[5]) if r[5] not in (None, '') else 80
head = '''/* ==========================================================================
   audio.js ― BGM・効果音の設定（「BGM・効果音シート」から tools/make_audio_js.py（xlsx）または tools/audio_from_dump.py（書き出したテキスト）で作る）
   ・BGM：場面ごとに流す曲。file が空＝前の曲をそのまま流す／'なし'＝止める。同じファイルが続くときは、曲を切らずに流し続ける
   ・SE ：効果音。file が空＝今の仮の音（プログラムで作った音）を鳴らす
   ・音のファイルは assets/audio/ に置くと、組み立て（build.js）のときにゲームの中へ埋め込まれる
   ========================================================================== */
window.AUDIO = '''
open(sys.argv[1], 'w', encoding='utf-8').write(head + json.dumps(cfg, ensure_ascii=False, indent=1) + ';\n')
print('wrote', sys.argv[1])
