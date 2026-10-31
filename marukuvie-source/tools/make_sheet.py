# セリフ修正用のスプレッドシートを作る（tools/export_lines.js の JSON から）
import json, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
d = json.load(open(sys.argv[1], encoding='utf-8')); out = sys.argv[2]
F = 'Arial'
HEAD = PatternFill('solid', fgColor='5E4D5A'); EDIT = PatternFill('solid', fgColor='FFF2A8'); PROV = PatternFill('solid', fgColor='ECE6F5'); ORIG = PatternFill('solid', fgColor='DFF0EA'); LOCK = PatternFill('solid', fgColor='F2F0EA')
thin = Side(style='thin', color='C9C0C6'); BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
wb = Workbook()

def table(ws, heads, widths, rows, edit_cols, wrap_cols, len_pairs):
    ws.append(heads)
    for c, wdt in enumerate(widths, 1):
        cell = ws.cell(1, c); cell.font = Font(name=F, bold=True, color='FFFFFF'); cell.fill = HEAD; cell.alignment = Alignment(vertical='center', wrap_text=True); cell.border = BOX
        ws.column_dimensions[cell.column_letter].width = wdt
    for r in rows: ws.append(r)
    n = len(rows) + 1
    for row in ws.iter_rows(min_row=2, max_row=n):
        for cell in row:
            cell.font = Font(name=F, size=10); cell.border = BOX
            cell.alignment = Alignment(vertical='top', wrap_text=cell.column in wrap_cols)
            if cell.column in edit_cols: cell.fill = EDIT
            elif cell.column not in wrap_cols: cell.fill = LOCK
    for src, dst in len_pairs:   # 文字数（数式）
        sl = ws.cell(1, src).column_letter
        for r in range(2, n + 1):
            c = ws.cell(r, dst); c.value = '=IF(%s%d="","",LEN(%s%d))' % (sl, r, sl, r); c.alignment = Alignment(vertical='top', horizontal='center'); c.font = Font(name=F, size=10, color='7A6B77')
    ws.freeze_panes = 'B2'; ws.auto_filter.ref = ws.dimensions; ws.row_dimensions[1].height = 34
    return n

# ---- 使い方 ----
ws = wb.active; ws.title = '使い方'
ws.column_dimensions['A'].width = 4; ws.column_dimensions['B'].width = 30; ws.column_dimensions['C'].width = 90
rows = [
 ('マルクヴィエ物語　セリフ修正シート', None),
 (None, None),
 ('修正のしかた', '黄色の列「修正後のセリフ」に、新しい文を書いてください。直さない行は空欄のままで大丈夫です（空欄＝今の文のまま）。'),
 ('', '書き終わったら、このファイルをそのまま送ってください。黄色の列の内容をゲームに反映します。'),
 ('書いてはいけない列', 'A列「No」とB列「ID」は、ゲームのどの文かを見分けるための番号です。変えたり、行を並べ替えて消したりしないでください（並べ替え・絞り込みはOK）。'),
 ('「現在のセリフ」の列', '今ゲームに入っている文です。ここを直しても反映されません（比べるための列です）。'),
 ('区分', '「原稿」＝ストーリーフロー（v1.1・v6.0）の文をそのまま使った文。「仮」＝原稿にない部分を試作のために仮に書いた文。差し替えたいのは主に「仮」の行です。'),
 ('文字数', '数式で自動計算されます。メッセージ枠には、パソコンで1行40文字前後・3行ほどが目安です。'),
 ('改行', 'セルの中で改行（Alt＋Enter）すると、ゲームでもその位置で改行されます。'),
 ('メモ', '「メモ」列には、気になることや、言い回しの相談などを自由に書いてください。'),
 ('シートの中身', '「セリフ」＝人物の台詞と地の文（%d行）／「案内文」＝ヒントや選べない理由など（%d行）／「結末」＝結末の名前と説明／「選択肢」＝選択肢のボタンの文（表 v6.0 の文そのまま。直すと表とずれるので注意）' % (len(d['lines']), len(d['guides']))),
 (None, None),
 ('記入例', None),
]
for i, (a, b) in enumerate(rows, 1):
    ws.cell(i, 2, a); ws.cell(i, 3, b)
    ws.cell(i, 2).font = Font(name=F, bold=True, size=16 if i == 1 else 11, color='3E333C'); ws.cell(i, 3).font = Font(name=F, size=11)
    ws.cell(i, 3).alignment = Alignment(wrap_text=True, vertical='top'); ws.cell(i, 2).alignment = Alignment(vertical='top', wrap_text=True)
ex = len(rows) + 1
for j, (h, v) in enumerate([('話し手', 'アルボ'), ('現在のセリフ', 'ルノ……すまんな、心配をかけて。少し休めば、よくなるさ。'), ('修正後のセリフ（例）', 'ルノ、すまんのう。なに、ひと晩眠れば元どおりじゃ。')]):
    ws.cell(ex + j, 2, h).font = Font(name=F, bold=True); c = ws.cell(ex + j, 3, v); c.font = Font(name=F, size=11)
    if j == 2: c.fill = EDIT
ws.cell(ex + 3, 3, '※記入例はこのシートだけの見本です。ゲームには反映されません。').font = Font(name=F, size=9, color='7A6B77')

# ---- セリフ ----
ws = wb.create_sheet('セリフ')
heads = ['No', 'ID（変えない）', '区画', '場面', '話し手', '区分', '出典', 'どこで出るか', '現在のセリフ', '修正後のセリフ（ここに入力）', '文字数（現在）', '文字数（修正後）', 'メモ', '照合（変えない）']
sec = lambda s: '序章' if s in ('区画0', '序章') else s
rows = [[i + 1, r['id'], sec(r['sec']), r['scene'], r['who'], r['kind'], r['src'], r['note'], r['text'], None, None, None, None, r['occ']] for i, r in enumerate(d['lines'])]
n = table(ws, heads, [5, 18, 8, 16, 9, 6, 7, 22, 52, 52, 8, 8, 24, 6], rows, {10, 13}, {8, 9, 10, 13}, [(9, 11), (10, 12)])
for r in range(2, n + 1):
    k = ws.cell(r, 6); k.fill = PROV if k.value == '仮' else ORIG
ws.column_dimensions['B'].hidden = False; ws.column_dimensions['N'].hidden = True

# ---- 案内文 ----
ws = wb.create_sheet('案内文')
WHERE = {'P.s1_restack': '区画1：「二人で協力して登る」で足場が崩れたあと', 'P.s2_revealed': '区画2：フラメトがスイッチを見つけたあと', 'P.s4_bonus': '区画4：隠し部屋が開いたとき', 'P.s9_why_k': '区画9・10：選べない選択肢の理由（カルマが高い）', 'P.s9_why_short': '区画9：選べない選択肢の理由（短縮ルート）',
         'P.s10_why_call': '区画10：「呼びかける」が選べない理由', 'P.s10_why_home': '区画10：「フラメトと帰る」がまだ選べない理由', 'P.s10_why_hurt': '区画10：フラメトを攻撃したあとの理由', 'P.s4_back_mirror': '区画4：扉の前の選択肢（謎解きに戻る）', 'P.s4_hint_secret': '区画4：隠し部屋が開いていないときの手がかり',
         'P.t_warn': '展示モード：残り5分のお知らせ', 'P.t_last': '展示モード：残り30秒のお知らせ', 'P.s4_leave_it': '区画4：隠し部屋の選択肢（何も取らない）'}
rows = [[i + 1, r['id'], sec(r['sec']), WHERE.get(r['id'], ''), r['text'], None, None, None, None, r['occ']] for i, r in enumerate(d['guides'])]
table(ws, ['No', 'ID（変えない）', '区画', 'どこで出るか', '現在の文', '修正後の文（ここに入力）', '文字数（現在）', '文字数（修正後）', 'メモ', '照合（変えない）'], [5, 18, 8, 36, 44, 44, 8, 8, 24, 6], rows, {6, 9}, {4, 5, 6, 9}, [(5, 7), (6, 8)])
ws.column_dimensions['J'].hidden = True

# ---- 結末 ----
ws = wb.create_sheet('結末')
rows = [[i + 1, e['id'], e['mark'], e['name'], None, e['note'], None, None] for i, e in enumerate(d['ends'])]
table(ws, ['No', 'ID（変えない）', '印', '現在の名前', '修正後の名前（ここに入力）', '現在の説明', '修正後の説明（ここに入力）', 'メモ'], [5, 18, 5, 24, 24, 44, 44, 24], rows, {5, 7, 8}, {4, 5, 6, 7, 8}, [])

# ---- 選択肢 ----
ws = wb.create_sheet('選択肢')
rows = [[i + 1, c['id'], sec(c['sec']), c['row'], c['label'], None, c['K'], c['A'], None] for i, c in enumerate(d['choices'])]
table(ws, ['No', 'ID（変えない）', '区画', '表の行（v6.0）', '現在のボタンの文', '修正後の文（ここに入力）', 'カルマ差', '不安差', 'メモ'], [5, 16, 8, 9, 44, 44, 8, 8, 24], rows, {6, 9}, {5, 6, 9}, [])
ws.cell(len(rows) + 3, 2, '※ボタンの文は、ストーリーフロー v6.0「選択肢と分岐」の表の文をそのまま使っています。直す場合は、表のほうもあわせて直すことをおすすめします。数値（カルマ差・不安差）はここでは変えられません。').font = Font(name=F, size=9, color='7A6B77')
wb.save(out); print('saved', out)
