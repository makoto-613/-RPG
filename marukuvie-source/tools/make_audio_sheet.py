# BGM・効果音の設定シートを作る
import sys, os, json
sys.path.insert(0, os.path.dirname(__file__)); from audio_def import BGM, SE
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
out = sys.argv[1]; cur = {}
if len(sys.argv) > 2:   # 今の audio.js の値を入れて作る
    t = open(sys.argv[2], encoding='utf-8').read(); cur = json.loads(t[t.index('{'):t.rindex('}') + 1])
F = 'Arial'; HEAD = PatternFill('solid', fgColor='5E4D5A'); EDIT = PatternFill('solid', fgColor='FFF2A8'); LOCK = PatternFill('solid', fgColor='F2F0EA')
thin = Side(style='thin', color='C9C0C6'); BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
wb = Workbook()
def table(ws, heads, widths, rows, edit):
    ws.append(heads)
    for c, w in enumerate(widths, 1):
        x = ws.cell(1, c); x.font = Font(name=F, bold=True, color='FFFFFF'); x.fill = HEAD; x.alignment = Alignment(vertical='center', wrap_text=True); x.border = BOX
        ws.column_dimensions[x.column_letter].width = w
    for r in rows: ws.append(r)
    for row in ws.iter_rows(min_row=2, max_row=len(rows) + 1):
        for x in row:
            x.font = Font(name=F, size=10); x.border = BOX; x.alignment = Alignment(vertical='top', wrap_text=True); x.fill = EDIT if x.column in edit else LOCK
    ws.freeze_panes = 'C2'; ws.row_dimensions[1].height = 34; ws.auto_filter.ref = ws.dimensions
    return len(rows) + 1
ws = wb.active; ws.title = '使い方'
ws.column_dimensions['A'].width = 4; ws.column_dimensions['B'].width = 26; ws.column_dimensions['C'].width = 92
L = [('マルクヴィエ物語　BGM・効果音シート', None), (None, None),
 ('書き方', '黄色の列に、使いたい音のファイル名（例：title_theme.mp3）と、音量などを書いてください。'),
 ('', '書き終わったら、このシートと、音のファイル（mp3 など）を一緒に送ってください。ゲームの中に埋め込みます。'),
 ('BGM：空欄のとき', '前の場面の曲を、そのまま流し続けます（曲が切れません）。'),
 ('BGM：「なし」と書いたとき', 'その場面で曲を止めます（無音にしたい場面に）。'),
 ('BGM：同じファイル名が続くとき', '場面が変わっても、曲を最初から流し直さず、そのまま流し続けます。'),
 ('効果音：空欄のとき', '今の仮の音（プログラムで作った簡単な音）が鳴ります。'),
 ('音量', '0〜100。BGMは60前後、効果音は80前後が目安です。'),
 ('ループ・フェード', 'ループ：「する」なら曲の終わりで最初に戻る。フェード：曲が切り替わるときに、前の曲が消えて次の曲が大きくなるまでの秒数。'),
 ('ファイルの形式', 'mp3 がおすすめです（ogg・wav・m4a も使えます）。ファイル名は半角英数字にすると安全です。'),
 ('ファイルの大きさ', 'ゲームは1つのファイルにまとめるため、音は合計で12MBくらいまでにしてください（BGMは1曲1〜2MBが目安。短めのループ曲が向いています）。'),
 ('使ってよい音か', 'フリー素材でも、利用規約（クレジット表記・加工の可否・公開の範囲）を必ず確認してください。'),
 ('書いてはいけない列', '「ID」の列は、ゲームのプログラムが使う名前です。変えないでください。'),
 (None, None), ('記入例', None)]
for i, (a, b) in enumerate(L, 1):
    ws.cell(i, 2, a).font = Font(name=F, bold=True, size=16 if i == 1 else 11, color='3E333C'); c = ws.cell(i, 3, b); c.font = Font(name=F, size=11); c.alignment = Alignment(wrap_text=True, vertical='top'); ws.cell(i, 2).alignment = Alignment(wrap_text=True, vertical='top')
ex = len(L) + 1
for j, (h, v) in enumerate([('どこで流れるか', 'タイトル画面'), ('曲ファイル名', 'title_theme.mp3'), ('音量／ループ／フェード', '60／する／1.5')]):
    ws.cell(ex + j, 2, h).font = Font(name=F, bold=True); c = ws.cell(ex + j, 3, v); c.font = Font(name=F, size=11); c.fill = EDIT if j else PatternFill()
ws.cell(ex + 3, 3, '※記入例はこのシートだけの見本です。').font = Font(name=F, size=9, color='7A6B77')
ws = wb.create_sheet('BGM')
b = cur.get('BGM', {})
rows = [[i + 1, k, w, (b.get(k, {}).get('file') or None), b.get(k, {}).get('vol', 60), 'する' if b.get(k, {}).get('loop', True) else 'しない', b.get(k, {}).get('fade', 1.0), None] for i, (k, w) in enumerate(BGM)]
n = table(ws, ['No', 'ID（変えない）', 'どこで流れるか', '曲ファイル名（ここに入力）', '音量（0〜100）', 'ループ', 'フェード（秒）', 'メモ'], [5, 13, 44, 32, 11, 9, 11, 30], rows, {4, 5, 6, 7, 8})
dv = DataValidation(type='list', formula1='"する,しない"', allow_blank=True); ws.add_data_validation(dv); dv.add('F2:F%d' % n)
dv2 = DataValidation(type='whole', operator='between', formula1='0', formula2='100', allow_blank=True); ws.add_data_validation(dv2); dv2.add('E2:E%d' % n)
ws = wb.create_sheet('効果音')
e = cur.get('SE', {})
rows = [[i + 1, k, w, now, (e.get(k, {}).get('file') or None), e.get(k, {}).get('vol', 80), None] for i, (k, w, now) in enumerate(SE)]
n = table(ws, ['No', 'ID（変えない）', 'どこで鳴るか', '今の音（仮の合成音）', '音ファイル名（ここに入力）', '音量（0〜100）', 'メモ'], [5, 11, 40, 24, 32, 11, 30], rows, {5, 6, 7})
dv3 = DataValidation(type='whole', operator='between', formula1='0', formula2='100', allow_blank=True); ws.add_data_validation(dv3); dv3.add('F2:F%d' % n)
wb.save(out); print('saved', out)
