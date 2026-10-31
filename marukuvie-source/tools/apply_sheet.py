# セリフ修正シート（xlsx）の「修正後」の列を、story.js・rules.js に書き戻す
# 使い方: python3 tools/apply_sheet.py 修正済み.xlsx [story.js] [rules.js]
import sys, json, re
from openpyxl import load_workbook
xl = sys.argv[1]; SP = sys.argv[2] if len(sys.argv) > 2 else 'story.js'; RP = sys.argv[3] if len(sys.argv) > 3 else 'rules.js'
story = open(SP, encoding='utf-8').read(); rules = open(RP, encoding='utf-8').read()
q1 = lambda t: "'" + t.replace('\\', '\\\\').replace("'", "\\'").replace('\r\n', '\n').replace('\n', '\\n') + "'"
log, err = [], []
def nth_replace(src, old, new, n, where):
    idx, start = -1, 0
    for _ in range(n):
        idx = src.find(old, start)
        if idx < 0: err.append('%s：元の文が見つかりません（%s）' % (where, old[:30])); return src
        start = idx + len(old)
    return src[:idx] + new + src[idx + len(old):]
wb = load_workbook(xl)
def rows(name, oc, nc, occ_c=None):
    ws = wb[name]; head = [c.value for c in ws[1]]
    for r in ws.iter_rows(min_row=2, values_only=True):
        if r[1] is None or r[nc] in (None, ''): continue
        new = str(r[nc]).replace('\r\n', '\n')
        if new == r[oc]: continue
        yield r, new
for sheet, oc, nc in (('セリフ', 8, 9), ('案内文', 4, 5)):
    for r, new in rows(sheet, oc, nc):
        story = nth_replace(story, q1(r[oc]), q1(new), int(r[-1] if sheet == 'セリフ' else r[9]), '%s No.%s' % (sheet, r[0])); log.append('%s No.%s: %s → %s' % (sheet, r[0], r[oc], new))
for r, new in rows('結末', 3, 4):
    story = nth_replace(story, "name: " + q1(r[3]), "name: " + q1(new), 1, '結末 No.%s' % r[0]); log.append('結末名 %s → %s' % (r[3], new))
for r, new in rows('結末', 5, 6):
    story = nth_replace(story, "note: " + q1(r[5]), "note: " + q1(new), 1, '結末 No.%s' % r[0]); log.append('結末説明 %s → %s' % (r[5], new))
for r, new in rows('選択肢', 4, 5):
    rules = nth_replace(rules, 'label: ' + json.dumps(r[4], ensure_ascii=False), 'label: ' + json.dumps(new, ensure_ascii=False), 1, '選択肢 No.%s' % r[0]); log.append('選択肢 %s → %s' % (r[4], new))
if err: print('\n'.join(err)); sys.exit(1)
open(SP, 'w', encoding='utf-8').write(story); open(RP, 'w', encoding='utf-8').write(rules)
print('反映した数: %d' % len(log)); print('\n'.join(log))
